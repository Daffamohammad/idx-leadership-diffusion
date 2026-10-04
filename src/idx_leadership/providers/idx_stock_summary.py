"""Official IDX closing observations, separate from adjusted-return analytics."""
from __future__ import annotations

import hashlib
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_stock_summary(path: Path) -> list[dict[str, Any]]:
    # The official-source validator already uses this optional workbook parser.
    # Core normalization remains independent of it for fresh-clone tests.
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("SOURCE_UNAVAILABLE: reading IDX workbooks requires openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        required = {"Stock Code", "Previous", "Close", "Last Trading Date", "Listed Shares"}
        if not required.issubset(header):
            raise ValueError(f"Stock Summary missing columns: {sorted(required.difference(header))}")
        return [dict(zip(header, row)) for row in rows if any(value is not None for value in row)]
    finally:
        workbook.close()


def _number(value: Any, *, positive: bool = False) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(result) or (positive and result <= 0):
        return None
    return result


def _trade_date(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    for pattern in ("%d %b %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value), pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Unrecognized official trade date: {value!r}")


def normalize_stock_summary(
    rows: Iterable[Mapping[str, Any]],
    *,
    as_of: str,
    classifications: Iterable[Mapping[str, Any]] = (),
) -> dict[str, Any]:
    date.fromisoformat(as_of)
    metadata = {str(row["ticker"]): row for row in classifications}
    records = []
    seen: set[str] = set()
    for raw in rows:
        code = str(raw.get("Stock Code") or "").strip().upper()
        if not code:
            continue
        if not code.isalnum() or not 4 <= len(code) <= 5:
            raise ValueError(f"Invalid IDX stock code: {code!r}")
        ticker = f"{code}.JK"
        if ticker in seen:
            raise ValueError(f"Duplicate official stock code: {code}")
        seen.add(ticker)
        trade_date = _trade_date(raw.get("Last Trading Date"))
        if trade_date > as_of:
            raise ValueError(f"Future trade date for {ticker}: {trade_date} > {as_of}")
        close = _number(raw.get("Close"), positive=True)
        previous = _number(raw.get("Previous"), positive=True)
        listed_shares = _number(raw.get("Listed Shares"), positive=True)
        observed = close is not None and previous is not None and trade_date == as_of
        reference = metadata.get(ticker, {})
        taxonomy = reference.get("taxonomy") or {}
        classification_date = reference.get("source_as_of")
        if classification_date and str(classification_date)[:10] > as_of:
            raise ValueError(f"Future classification for {ticker}")
        volume = _number(raw.get("Volume"))
        if volume is not None and volume < 0:
            raise ValueError(f"Negative volume for {ticker}")
        foreign_buy = _number(raw.get("Foreign Buy"))
        foreign_sell = _number(raw.get("Foreign Sell"))
        if any(v is not None and v < 0 for v in (foreign_buy, foreign_sell)):
            raise ValueError(f"Negative foreign share quantity for {ticker}")
        company_name = str(raw.get("Company Name") or code)
        multiple_voting = company_name.upper().startswith("MVS ")
        records.append({
            "ticker": ticker,
            "company_name": company_name,
            "instrument_type": "MULTIPLE_VOTING_SHARES" if multiple_voting else "LISTED_STOCK",
            "analysis_requested": not multiple_voting,
            "as_of": as_of,
            "last_trade_date": trade_date,
            "close": close,
            "previous_close": previous,
            "return_1d": (close / previous - 1) * 100 if observed else None,
            "listed_shares": listed_shares,
            "market_cap": close * listed_shares if close and listed_shares else None,
            "weight_for_index": _number(raw.get("Weight For Index"), positive=True),
            "volume_shares": volume,
            "traded": not multiple_voting and observed and volume is not None and volume > 0,
            "value_idr": _number(raw.get("Value")),
            "foreign_buy_shares": foreign_buy,
            "foreign_sell_shares": foreign_sell,
            "foreign_net_shares": foreign_buy - foreign_sell if foreign_buy is not None and foreign_sell is not None else None,
            "taxonomy": taxonomy,
            "classification_source": reference.get("source"),
            "classification_as_of": classification_date,
            "listing_board": reference.get("listing_board"),
        })
    if not records or not any(r["last_trade_date"] == as_of for r in records):
        raise ValueError("No official observations for the requested session")
    records.sort(key=lambda row: row["ticker"])
    traded = [r for r in records if r["traded"]]
    return {
        "schema_version": "market-daily-v1",
        "as_of": as_of,
        "price_basis": "official_close",
        "units": {"price": "IDR per share", "market_cap": "IDR", "foreign": "shares"},
        "listed_count": len(records),
        "observed_price_count": sum(r["return_1d"] is not None for r in records),
        "classification_count": sum(bool(r["taxonomy"].get("sector")) for r in records),
        "breadth": {
            "scope": "Stocks with positive volume and comparable official closes",
            "advancers": sum(r["return_1d"] > 0 for r in traded),
            "flat": sum(r["return_1d"] == 0 for r in traded),
            "decliners": sum(r["return_1d"] < 0 for r in traded),
            "traded_count": len(traded),
            "not_traded_or_unavailable": len(records) - len(traded),
        },
        "records": records,
    }
